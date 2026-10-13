"""Помощники тестов: docx с полями подстановки в разных частях пакета."""
import base64
import io
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _run(text):
    r = etree.SubElement(etree.Element("dummy"), qn("w:r"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return r


def make_docx(body_split=None, header=None, table=None, hyperlink=None, ins=None, sdt=None):
    """docx с полями в разных местах пакета.

    body_split: список кусков одного поля, разложенных по прогонам
    hyperlink / ins / sdt: текст, обёрнутый в w:hyperlink / w:ins / w:sdt
    """
    doc = Document()
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    if body_split:
        p = doc.add_paragraph()
        for chunk in body_split:
            p.add_run(chunk)
    if table:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = table
    for tag, text in (("w:hyperlink", hyperlink), ("w:ins", ins)):
        if text:
            p = doc.add_paragraph()
            wrapper = etree.SubElement(p._p, qn(tag))
            if tag == "w:ins":
                wrapper.set(qn("w:id"), "1")
                wrapper.set(qn("w:author"), "t")
            wrapper.append(_run(text))
    if sdt:
        p = doc.add_paragraph()
        s = etree.SubElement(p._p, qn("w:sdt"))
        etree.SubElement(s, qn("w:sdtPr"))
        content = etree.SubElement(s, qn("w:sdtContent"))
        content.append(_run(sdt))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_text(raw):
    """Весь текст всех частей пакета (тело + колонтитулы)."""
    z = zipfile.ZipFile(io.BytesIO(raw))
    out = []
    for name in z.namelist():
        if name.startswith("word/") and name.endswith(".xml"):
            tree = etree.fromstring(z.read(name))
            out.append("".join(t.text or "" for t in tree.iter("{%s}t" % W)))
    return "\n".join(out)


def docx_parts(raw):
    return sorted(zipfile.ZipFile(io.BytesIO(raw)).namelist())


def b64(raw):
    return base64.b64encode(raw)


def _u(record):
    """Перед сменой пользователя сбрасываем кэш: в проде каждый запрос —
    один пользователь, а в тесте вычисляемые поля без depends_context('uid')
    иначе переносят значения от одного пользователя к другому."""
    record.env.invalidate_all()
    return record


def docx_t_nodes(raw):
    """[(текст, есть ли xml:space=preserve)] для всех w:t основного документа."""
    z = zipfile.ZipFile(io.BytesIO(raw))
    tree = etree.fromstring(z.read("word/document.xml"))
    space = "{http://www.w3.org/XML/1998/namespace}space"
    return [(t.text or "", t.get(space) == "preserve") for t in tree.iter("{%s}t" % W)]
