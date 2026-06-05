from odoo import _, models, Command, fields, api
from datetime import datetime
from odoo.exceptions import ValidationError
import base64
import io
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image
from docx.shared import Pt
import tempfile
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os
from PyPDF2 import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from datetime import timedelta
from docx.shared import Mm
from reportlab.lib.pagesizes import A4
from lxml import etree

import logging

_logger = logging.getLogger(__name__)


class CorrOutgoingApproveProcessMixin(models.AbstractModel):
    _name = "corr.outgoing.approve.process.mixin"
    _description = "Correspondence Outgoing approve process mixin"

    number_of_esp_signs = fields.Integer(
        string="Количество ЭЦП подписей",
        default=0,
        copy=False,
    )

    # ==================================================================
    # Маршрутизация статусов
    # ==================================================================

    def _get_next_state_after(self, current_state):
        """
        Возвращает следующий статус после текущего по workflow.

        Стандартный тип:
            under_approval → approval → processing → review → done

        medical_examination:
            under_approval → approval_medical_examination → processing → review → done
        """
        is_med = (
            hasattr(self, '_is_medical_examination_type')
            and self._is_medical_examination_type()
        )

        state_sequence = {
            'under_approval': 'approval_medical_examination' if is_med else 'approval',
            'approval': 'processing',
            'approval_medical_examination': 'processing',
            'processing': 'review',
            'review': 'done',
        }
        return state_sequence.get(current_state)

    # ==================================================================
    # Создание линий согласования
    # ==================================================================

    def get_agreement_lines(self, state="under_approval"):
        """Создаёт линии согласования для указанного статуса."""
        # Защита от повторного создания если текущий user уже in_progress
        current_user_line = self.state_agreement_line_ids.filtered(
            lambda l: l.user_id.id == self.env.uid and l.status == 'in_progress'
        )
        if current_user_line:
            return True

        if hasattr(self, "method_on_start"):
            self.method_on_start()

        # Получаем need_esp из workflow
        workflow = self.env['appstream.approval.workflow'].sudo().search([
            ('approval_id.model_id.model', '=', self._name),
            ('state', '=', state)
        ], limit=1)
        need_esp = workflow.need_esp if workflow else False

        # Очищаем текущие линии
        self.sudo().state_agreement_line_ids.unlink()

        # Создаём линии из before_approval_groups (модельная логика)
        before_lines, sequence, init_approvers = (
            self.get_agreement_lines_by_before_approval_groups(
                status=state, need_esp=need_esp
            )
        )
        agreement_lines = before_lines
        start = sequence + 1 if before_lines else sequence

        # Создаём линии из approval_group_ids workflow (XML-конфигурация)
        approval_lines = self.get_agreement_lines_by_approval_groups(
            start=start, init_approvers=init_approvers, need_esp=need_esp
        )
        agreement_lines += approval_lines

        # Если линии пустые - пропускаем статус
        if not agreement_lines:
            next_state = self._get_next_state_after(state)
            if next_state and next_state != 'done':
                return self.get_agreement_lines(next_state)
            elif next_state == 'done':
                self.action_notify("approved")
                return self.sudo().write({"state": "done"})
            return self.sudo().write({"state": state})

        self.sudo().state_agreement_line_ids = agreement_lines

        if hasattr(self, "check_agreement_line_status"):
            self.check_agreement_line_status()

        return self.sudo().write({"state": state})

    def get_agreement_lines_by_before_approval_groups(
        self, agreement_lines=None, sequence=1, init_approvers=None,
        status=None, need_esp=False
    ):
        """
        Создаёт линии согласования в зависимости от статуса.
        Это основная логика маршрутизации людей по статусам.
        """
        agreement_lines = [] if agreement_lines is None else agreement_lines
        init_approvers = [] if init_approvers is None else init_approvers
        approver_ids = []

        # ==============================================================
        # under_approval — руководитель + дополнительные согласующие
        # ==============================================================
        if status == 'under_approval':
            initiator = self.create_uid
            if initiator and initiator.employee_id:
                manager = initiator.employee_id.parent_id
                if manager and manager.user_id:
                    manager_user = manager.user_id
                    is_director = manager_user.has_group(
                        'correspondence.group_correspondence_director'
                    )
                    if not is_director and manager_user.id != 2:
                        approver_ids.append(manager_user)

        # ==============================================================
        # review — инициатор (ознакомление)
        # ==============================================================
        if status == 'review':
            approver_ids.append(self.create_uid)

        # ==============================================================
        # approval — esp_signer_id подписывает с ЭЦП (стандартные типы)
        #   для medical_checkup: + employee_id (ЭЦП в системе, сотрудник)
        # ==============================================================
        if status == 'approval':
            if hasattr(self, 'esp_signer_id') and self.esp_signer_id:
                signer = self.esp_signer_id
                if signer.id != 2 and signer not in init_approvers:
                    init_approvers.append(signer)
                    self.action_notify("agreement", signer)
                    agreement_lines.append(
                        Command.create({
                            "sequence": sequence,
                            "model": self._name,
                            "user_id": signer.id,
                            "status": "in_progress",
                            "need_esp": True,
                        })
                    )
                    sequence += 1

            # Для мед. осмотра и объяснительной: employee_id подписывает после утверждающего
            if (
                hasattr(self, '_is_medical_checkup_type')
                and (self._is_medical_checkup_type() or self._is_explanation_type())
                and hasattr(self, 'employee_id')
                and self.employee_id
                and self.employee_id.user_id
            ):
                emp_user = self.employee_id.user_id
                if emp_user.id != 2 and emp_user not in init_approvers:
                    init_approvers.append(emp_user)
                    agreement_lines.append(
                        Command.create({
                            "sequence": sequence,
                            "model": self._name,
                            "user_id": emp_user.id,
                            "status": "waiting",
                            "need_esp": True,
                            "all_approve": True,
                        })
                    )

            return agreement_lines, sequence, init_approvers

        # ==============================================================
        # approval_medical_examination — 3 подписанта с ЭЦП
        #   seq 1: esp_signer_id        (ЭЦП в системе, утверждающий)
        #   seq 2: medical_worker_id    (ЭЦП на портале, мед. работник)
        #   seq 3: employee_id.user_id  (ЭЦП в системе, сотрудник)
        # ==============================================================
        if status == 'approval_medical_examination':
            return self._build_approval_medical_examination_lines()

        # ==============================================================
        # processing — секретарь
        # ==============================================================
        if status == 'processing':
            secretary_users = self.env['res.users'].sudo().search([
                ('groups_id', 'in', self.env.ref(
                    'correspondence.group_correspondence_secretary'
                ).id),
                ('id', '!=', 2),
            ])
            for secretary in secretary_users:
                if secretary not in init_approvers:
                    init_approvers.append(secretary)
                    self.action_notify("agreement", secretary)
                    agreement_lines.append(
                        Command.create({
                            "sequence": sequence,
                            "model": self._name,
                            "user_id": secretary.id,
                            "status": "in_progress",
                            "need_esp": need_esp,
                        })
                    )
                    break  # Один секретарь
            return agreement_lines, sequence, init_approvers

        # ==============================================================
        # Стандартная логика для approver_ids (under_approval, review)
        # ==============================================================
        for approver_id in approver_ids:
            init_approvers.append(approver_id)
            self.action_notify("agreement", approver_id)
            agreement_lines.append(
                Command.create({
                    "sequence": sequence,
                    "model": self._name,
                    "user_id": approver_id.id,
                    "status": "in_progress" if sequence == 1 else "waiting",
                    "need_esp": need_esp,
                })
            )

        # Дополнительные согласующие для under_approval
        if (
            status == 'under_approval'
            and hasattr(self, 'additional_approver_ids')
            and self.additional_approver_ids
        ):
            sequence += 1 if approver_ids else 0
            for user in self.additional_approver_ids:
                if user.id in [u.id for u in init_approvers] or user.id == 2:
                    continue
                init_approvers.append(user)
                agreement_lines.append(
                    Command.create({
                        "sequence": sequence,
                        "model": self._name,
                        "user_id": user.id,
                        "status": (
                            "in_progress"
                            if sequence == 1 and not approver_ids
                            else "waiting"
                        ),
                        "all_approve": True,
                        "need_esp": need_esp,
                    })
                )

        return agreement_lines, sequence, init_approvers

    def _build_approval_medical_examination_lines(self):
        """
        Создаёт 3 линии подписания с ЭЦП для medical_examination:
            seq 1 — esp_signer_id        (ЭЦП в системе, утверждающий)
            seq 2 — medical_worker_id    (ЭЦП на портале, мед. работник)
            seq 3 — employee_id.user_id  (ЭЦП в системе, сотрудник)
        """
        agreement_lines = []
        init_approvers = []
        sequence = 1

        # 1. Утверждающий (esp_signer_id) — ЭЦП в системе
        if hasattr(self, 'esp_signer_id') and self.esp_signer_id:
            signer = self.esp_signer_id
            if signer.id != 2 and signer not in init_approvers:
                init_approvers.append(signer)
                self.action_notify("agreement", signer)
                agreement_lines.append(
                    Command.create({
                        "sequence": sequence,
                        "model": self._name,
                        "user_id": signer.id,
                        "status": "in_progress",
                        "need_esp": True,
                        "all_approve": True,
                    })
                )
                sequence += 1

        # 2. Медицинский работник (medical_worker_id) — ЭЦП на портале
        if hasattr(self, 'medical_worker_id') and self.medical_worker_id:
            partner = self.medical_worker_id
            user = self.env['res.users'].sudo().search([
                ('partner_id', '=', partner.id)
            ], limit=1)
            if user and user.id != 2 and user not in init_approvers:
                init_approvers.append(user)
                agreement_lines.append(
                    Command.create({
                        "sequence": sequence,
                        "model": self._name,
                        "user_id": user.id,
                        "status": "waiting" if agreement_lines else "in_progress",
                        "need_esp": True,
                        "all_approve": True,
                    })
                )
                sequence += 1

        # 3. Сотрудник (employee_id) — ЭЦП в системе
        if (
            hasattr(self, 'employee_id')
            and self.employee_id
            and self.employee_id.user_id
        ):
            emp_user = self.employee_id.user_id
            if emp_user.id != 2 and emp_user not in init_approvers:
                init_approvers.append(emp_user)
                agreement_lines.append(
                    Command.create({
                        "sequence": sequence,
                        "model": self._name,
                        "user_id": emp_user.id,
                        "status": "waiting",
                        "need_esp": True,
                        "all_approve": True,
                    })
                )

        return agreement_lines, sequence, init_approvers

    # ==================================================================
    # Линии из approval_group_ids workflow
    # ==================================================================

    def get_agreement_lines_by_approval_groups(
        self, agreement_lines=None, start=1, init_approvers=None, need_esp=False
    ):
        """Получает согласующих из approval_group_ids workflow."""
        agreement_lines = [] if agreement_lines is None else agreement_lines
        init_approvers = [] if init_approvers is None else init_approvers
        approval_groups = self.state_id.approval_group_ids
        approval_groups = self.filter_connection(approval_groups)
        approval_groups = approval_groups.sorted("sequence")
        seen = set()
        sequence = start
        for approval_group in approval_groups:
            approval_group_users = self.check_group(approval_group.group_ids)
            users = approval_group.user_ids or approval_group_users
            users = users.filtered(
                lambda u: u not in init_approvers and u.id not in seen and u.id != 2
            )
            seen.update(users.ids)
            flag = False
            for user_id in users:
                if sequence == 1:
                    self.action_notify("agreement", user_id)
                flag = True
                agreement_lines.append(
                    Command.create({
                        "sequence": sequence,
                        "model": self._name,
                        "user_id": user_id.id,
                        "status": "in_progress" if sequence == 1 else "waiting",
                        "all_approve": approval_group.all_approve,
                        "need_esp": approval_group.need_esp or need_esp,
                    })
                )
            if flag:
                sequence += 1
        return agreement_lines

    # ==================================================================
    # Согласование: after_script + _process_post_approval
    # ==================================================================

    def get_current_coordinator(self):
        return self.state_agreement_line_ids.filtered(
            lambda line: line.user_id == self.env.user and line.status == "in_progress"
        )

    def after_script(self, next_state, cur_state):
        """
        Обрабатывает согласование через систему (кнопки в UI).
        Вызывается из after_script в XML workflow.

        Для портального подписания используется portal_sign_esp → _process_post_approval
        напрямую (approver уже помечен как agreed в portal_sign_esp).
        """
        current_coordinator = self.get_current_coordinator()
        if hasattr(self, "additional_condition"):
            self.additional_condition()
        if current_coordinator:
            # Помечаем как agreed
            current_coordinator.sudo().status = "agreed"
            current_coordinator.sudo().agreement_date = datetime.now()

            # Записываем в историю (+ вшивание ЭЦП в файл если signed)
            state_description = {
                sd[0]: sd[1]
                for sd in self._fields['state']._description_selection(self.env)
            }
            self.add_to_history(
                current_coordinator, state_description.get(cur_state)
            )

            # Удаляем activity текущего пользователя
            self._remove_approval_activity(user_id=self.env.uid)

            # Единая логика перехода
            self._process_post_approval(
                current_coordinator,
                next_state=next_state,
                cur_state=cur_state,
            )
        else:
            raise ValidationError(_("Вы не являетесь текущим согласующим"))

    def _process_post_approval(self, coordinator, next_state=None, cur_state=None, from_portal=False):
        """
        Единая логика после согласования — используется из:
        - after_script (системное подписание кнопками/ЭЦП)  → from_portal=False
        - portal_sign_esp (портальное подписание ЭЦП)       → from_portal=True

        Алгоритм:
        1. Удаляет параллельных (не all_approve) при той же sequence
        2. Если ещё есть in_progress — остаёмся на cur_state
        3. Активирует следующих по sequence (waiting → in_progress)
        4. Если все согласовали — переходим на next_state через get_agreement_lines

        ВАЖНО: _schedule_approval_activity вызывается ТОЛЬКО при from_portal=True.
        При системном пути (from_portal=False) activity планируется фреймворком
        appstream_approval в _action_approve после выполнения after_script,
        чтобы избежать дублирования активностей.
        """
        if cur_state is None:
            cur_state = self.state
        if next_state is None:
            next_state = self._get_next_state_after(cur_state)

        # 1. Удаляем параллельных без all_approve
        approvers_to_remove = self.state_agreement_line_ids.filtered(
            lambda line: line.status == "in_progress"
            and line.id != coordinator.id
            and not line.all_approve
        )
        if approvers_to_remove:
            approvers_to_remove.sudo().unlink()

        # 2. Если ещё есть in_progress — остаёмся на текущем статусе
        still_in_progress = self.state_agreement_line_ids.filtered(
            lambda line: line.status == "in_progress" and line.id != coordinator.id
        )
        if still_in_progress:
            self.sudo().write({"state": cur_state})
            return

        # 3. Ищем следующих по sequence (waiting)
        next_sequences = [
            line.sequence for line in self.state_agreement_line_ids
            if line.sequence > coordinator.sequence
        ]
        next_sequence = min(next_sequences) if next_sequences else -1

        next_coordinators = self.state_agreement_line_ids.filtered(
            lambda line: line.status == "waiting"
            and line.id != coordinator.id
            and line.sequence == next_sequence
        )

        if next_coordinators:
            # Активируем следующих подписантов
            for nc in next_coordinators:
                nc.sudo().status = "in_progress"
                self.action_notify("agreement", nc.user_id)

            # Планируем activity ТОЛЬКО при портальном подписании.
            # При системном пути activity планируется фреймворком
            # (_action_approve → schedule_activity после after_script).
            if from_portal:
                self._schedule_approval_activity(
                    users=next_coordinators.mapped('user_id')
                )

            self.sudo().write({"state": cur_state})
        else:
            # 4. Все согласовали — переход на следующий статус
            if next_state == "done":
                self.action_notify("approved")
                self.sudo().write({"state": next_state})
                if hasattr(self, "method_on_approved"):
                    self.method_on_approved()
            elif next_state:
                if hasattr(self, "method_in_middle"):
                    self.method_in_middle()
                self.get_agreement_lines(next_state)
            else:
                self.action_notify("approved")
                self.sudo().write({"state": "done"})

    # ==================================================================
    # История + вшивание ЭЦП в файл
    # ==================================================================

    @api.model
    def add_to_history(self, current_coordinator, state=False, status="Согласовано"):
        """Добавляет запись в историю согласования и вшивает ЭЦП в файл."""
        if hasattr(self, "state_agreement_history_line_ids"):
            new_status = status
            if state:
                if status in ["Согласовано", "Отклонено", "Уведомлено об ошибке"]:
                    new_status += " на статусе '" + state + "'"
                else:
                    new_status += " со статуса '" + state + "'"
            self.sudo().state_agreement_history_line_ids = [
                Command.create({
                    "model": self._name,
                    "user_id": current_coordinator.user_id.id,
                    "status": new_status,
                    "commentary": current_coordinator.commentary,
                    "agreement_date": current_coordinator.agreement_date,
                    "date": current_coordinator.date,
                    "certificate_start_date": current_coordinator.certificate_start_date,
                    "certificate_end_date": current_coordinator.certificate_end_date,
                    "serial_number": current_coordinator.serial_number,
                    "issued_by": current_coordinator.issued_by,
                    "issued_to": current_coordinator.issued_to,
                    "certificate_status": current_coordinator.certificate_status,
                    "signed": current_coordinator.signed,
                    "qr": current_coordinator.qr,
                })
            ]


        # ===============================================================
        # Брендирование файлов при ЭЦП подписании
        # ===============================================================
        if current_coordinator.signed:
            line = current_coordinator

            # 1. Дополнительные документы — ВСЕГДА, без номера документа
            if hasattr(self, 'attachment_additional_sign_ids') and self.attachment_additional_sign_ids:
                for attachment in self.attachment_additional_sign_ids:
                    try:
                        if attachment.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                            self._process_docx_signature(attachment, line, include_doc_number=False, field_name='attachment_additional_sign_ids')
                        elif attachment.mimetype == 'application/pdf':
                            self._process_pdf_signature(attachment, line, include_doc_number=False, field_name='attachment_additional_sign_ids')
                    except Exception as e:
                        _logger.error("Ошибка брендирования доп.файла %s: %s", attachment.name, str(e))

            # 2. Исходящее письмо — только типы БЕЗ печатной формы
            #    (explanation, medical_examination, medical_checkup — QR через docxtpl шаблон)
            skip_main_branding = False
            if hasattr(self, '_is_explanation_type') and self._is_explanation_type():
                skip_main_branding = True
            if hasattr(self, '_is_medical_examination_type') and self._is_medical_examination_type():
                skip_main_branding = True
            if hasattr(self, '_is_medical_checkup_type') and self._is_medical_checkup_type():
                skip_main_branding = True
            if hasattr(self, '_is_job_offer_type') and self._is_job_offer_type():
                skip_main_branding = True
            if hasattr(self, '_is_reference_type') and self._is_reference_type():
                skip_main_branding = True

            if not skip_main_branding and hasattr(self, 'attachment_to_sign_ids') and self.attachment_to_sign_ids:
                for attachment in self.attachment_to_sign_ids:
                    try:
                        if attachment.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                            self._process_docx_signature(attachment, line, include_doc_number=True)
                        elif attachment.mimetype == 'application/pdf':
                            self._process_pdf_signature(attachment, line, include_doc_number=True)
                    except Exception as e:
                        _logger.error("Ошибка брендирования файла %s: %s", attachment.name, str(e))

            self.sudo().number_of_esp_signs += 1

    # ==================================================================
    # Брендирование файлов (DOCX / PDF) — восстановленная рабочая версия
    # ==================================================================

    def _process_docx_signature(self, attachment, line, include_doc_number=True, field_name='attachment_to_sign_ids'):
        """Обрабатывает DOCX файл — добавляет боковую подпись ЭЦП на все страницы + страницу сертификата."""
        docx_io = io.BytesIO(base64.b64decode(attachment.datas))
        document = Document(docx_io)

        # Боковая подпись через header (появится на всех страницах)
        self._add_sidebar_signature_docx(document, line, include_doc_number=include_doc_number)

        # Страница "ДОКУМЕНТ ПОДПИСАН" с данными сертификата
        self._add_certificate_page_docx(document, line)

        # Сохраняем
        edited_docx_io = io.BytesIO()
        document.save(edited_docx_io)
        edited_docx_io.seek(0)
        edited_docx_base64 = base64.b64encode(edited_docx_io.read())

        new_attachment = self.env['ir.attachment'].sudo().create({
            'name': 'Подписан_' + attachment.name,
            'datas': edited_docx_base64,
            'mimetype': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'res_model': self._name,
            'res_id': self.id,
        })

        setattr(self.sudo(), field_name, [(3, attachment.id)])
        setattr(self.sudo(), field_name, [(4, new_attachment.id)])

    def _add_sidebar_signature_docx(self, document, line, include_doc_number=True):
        """
        Добавляет боковую подпись ЭЦП в DOCX через header (появится на всех страницах).
        QR-код + текст повёрнуты на 90° и размещены на левом поле.
        """
        if not line.qr:
            return

        # Сохраняем QR во временный файл
        qr_binary = base64.b64decode(line.qr)
        with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp:
            tmp.write(qr_binary)
            tmp.flush()
            qr_path = tmp.name

        # Формируем текст подписи
        signer_name = line.user_id.name or ''
        signature_text = f"{signer_name} подписал(а)"

        sign_date = line.agreement_date.strftime('%d.%m.%Y') if line.agreement_date else ''
        if include_doc_number:
            doc_number = self.name or ''
            doc_info_text = f"{doc_number} от {sign_date}"
        else:
            doc_info_text = f"от {sign_date}"

        # Работаем с header для размещения на всех страницах
        section = document.sections[0]
        header = section.header

        if not header.paragraphs:
            paragraph = header.add_paragraph()
        else:
            paragraph = header.paragraphs[0]

        run = paragraph.add_run()

        # Добавляем QR как inline чтобы получить relationship ID
        run.add_picture(qr_path, width=Mm(15))

        # Получаем rId изображения
        drawing = run._r.find(qn('w:drawing'))
        inline_elem = drawing.find(qn('wp:inline'))
        blip = inline_elem.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        r_embed = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')

        # Размеры и позиция
        qr_size = int(Mm(15))
        pos_x = int(Mm(6))     # от левого края страницы
        pos_y = int(Mm(270))   # от верха (внизу страницы)
        rotation = 270 * 60000  # поворот на 270° (снизу вверх)

        # Floating anchor для QR
        anchor_xml = f'''<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                    xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"
                    xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
                    distT="0" distB="0" distL="114300" distR="114300"
                    simplePos="0" relativeHeight="251660288" behindDoc="0"
                    locked="0" layoutInCell="1" allowOverlap="1">
            <wp:simplePos x="0" y="0"/>
            <wp:positionH relativeFrom="page">
                <wp:posOffset>{pos_x}</wp:posOffset>
            </wp:positionH>
            <wp:positionV relativeFrom="page">
                <wp:posOffset>{pos_y}</wp:posOffset>
            </wp:positionV>
            <wp:extent cx="{qr_size}" cy="{qr_size}"/>
            <wp:effectExtent l="0" t="0" r="0" b="0"/>
            <wp:wrapNone/>
            <wp:docPr id="100" name="QR Signature"/>
            <a:graphic>
                <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
                    <pic:pic>
                        <pic:nvPicPr>
                            <pic:cNvPr id="0" name="QR"/>
                            <pic:cNvPicPr/>
                        </pic:nvPicPr>
                        <pic:blipFill>
                            <a:blip r:embed="{r_embed}"/>
                            <a:stretch><a:fillRect/></a:stretch>
                        </pic:blipFill>
                        <pic:spPr>
                            <a:xfrm rot="{rotation}">
                                <a:off x="0" y="0"/>
                                <a:ext cx="{qr_size}" cy="{qr_size}"/>
                            </a:xfrm>
                            <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                        </pic:spPr>
                    </pic:pic>
                </a:graphicData>
            </a:graphic>
        </wp:anchor>'''

        # Заменяем inline на anchor
        anchor_elem = etree.fromstring(anchor_xml)
        drawing.clear()
        drawing.append(anchor_elem)

        # Добавляем текстовый бокс с подписью
        text_pos_x = pos_x - int(Mm(32))
        text_pos_y = pos_y - qr_size - int(Mm(30))
        text_width = int(Mm(80))
        text_height = int(Mm(12))

        textbox_xml = f'''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                          xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                          xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
            <wp:anchor distT="0" distB="0" distL="114300" distR="114300"
                       simplePos="0" relativeHeight="251661312" behindDoc="0"
                       locked="0" layoutInCell="1" allowOverlap="1">
                <wp:simplePos x="0" y="0"/>
                <wp:positionH relativeFrom="page">
                    <wp:posOffset>{text_pos_x}</wp:posOffset>
                </wp:positionH>
                <wp:positionV relativeFrom="page">
                    <wp:posOffset>{text_pos_y}</wp:posOffset>
                </wp:positionV>
                <wp:extent cx="{text_width}" cy="{text_height}"/>
                <wp:effectExtent l="0" t="0" r="0" b="0"/>
                <wp:wrapNone/>
                <wp:docPr id="101" name="Text Signature"/>
                <a:graphic>
                    <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                        <wps:wsp>
                            <wps:cNvSpPr txBox="1"/>
                            <wps:spPr>
                                <a:xfrm rot="{rotation}">
                                    <a:off x="0" y="0"/>
                                    <a:ext cx="{text_width}" cy="{text_height}"/>
                                </a:xfrm>
                                <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                                <a:noFill/>
                                <a:ln><a:noFill/></a:ln>
                            </wps:spPr>
                            <wps:txbx>
                                <w:txbxContent>
                                    <w:p>
                                        <w:pPr><w:jc w:val="left"/></w:pPr>
                                        <w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t>{signature_text}</w:t></w:r>
                                    </w:p>
                                    <w:p>
                                        <w:pPr><w:jc w:val="left"/></w:pPr>
                                        <w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t>{doc_info_text}</w:t></w:r>
                                    </w:p>
                                </w:txbxContent>
                            </wps:txbx>
                            <wps:bodyPr rot="0" vert="horz" wrap="square" anchor="t"/>
                        </wps:wsp>
                    </a:graphicData>
                </a:graphic>
            </wp:anchor>
        </w:drawing>'''

        text_run = paragraph.add_run()
        text_drawing_elem = etree.fromstring(textbox_xml)
        text_run._r.append(text_drawing_elem)

    def _add_certificate_page_docx(self, document, line):
        """Добавляет страницу с данными сертификата ЭЦП."""
        document.add_section()

        title = document.add_paragraph()
        run = title.add_run("ДОКУМЕНТ ПОДПИСАН")
        run.bold = True
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        p = document.add_paragraph()
        p.add_run(line.user_id.name or '')

        p = document.add_paragraph()
        p.add_run("Дата выдачи сертификата: ").bold = True
        p.add_run(str(line.certificate_start_date or ''))

        p = document.add_paragraph()
        p.add_run("Дата окончания сертификата: ").bold = True
        p.add_run(str(line.certificate_end_date or ''))

        p = document.add_paragraph()
        p.add_run("Серийный номер: ").bold = True
        p.add_run(str(line.serial_number or ''))

        p = document.add_paragraph()
        p.add_run("Кем выдан: ").bold = True
        p.add_run(str(line.issued_by or ''))

        p = document.add_paragraph()
        p.add_run("Кому выдан: ").bold = True
        p.add_run(str(line.issued_to or ''))

        p = document.add_paragraph()
        p.add_run("Статус сертификата: ").bold = True
        p.add_run(str(line.certificate_status or ''))

        p = document.add_paragraph()
        p.add_run("Дата подписи: ").bold = True
        sign_date = line.agreement_date.strftime('%d.%m.%Y') if line.agreement_date else ''
        p.add_run(sign_date)

        if line.qr:
            qr_binary_data = base64.b64decode(line.qr)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp_file:
                tmp_file.write(qr_binary_data)
                tmp_file.flush()
                document.add_picture(tmp_file.name, width=Inches(1.5))

    def _process_pdf_signature(self, attachment, line, include_doc_number=True, field_name='attachment_to_sign_ids'):
        """Обрабатывает PDF файл — добавляет боковую подпись ЭЦП на ВСЕ страницы + страницу сертификата."""
        pdf_io = io.BytesIO(base64.b64decode(attachment.datas))
        reader = PdfReader(pdf_io)
        writer = PdfWriter()

        # Регистрируем шрифты
        module_path = os.path.dirname(os.path.dirname(__file__))
        font_path = os.path.join(module_path, 'static', 'src', 'ttf', 'DejaVuSans.ttf')
        bold_font_path = os.path.join(module_path, 'static', 'src', 'ttf', 'DejaVuSans-Bold.ttf')

        try:
            pdfmetrics.registerFont(TTFont('DejaVuSans', font_path))
            pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', bold_font_path))
            font_name = 'DejaVuSans'
            font_name_bold = 'DejaVuSans-Bold'
        except Exception:
            font_name = 'Helvetica'
            font_name_bold = 'Helvetica-Bold'

        signer_name = line.user_id.name or ''
        sign_date = line.agreement_date.strftime('%d.%m.%Y') if line.agreement_date else ''

        if include_doc_number:
            doc_number = self.name or ''
            doc_info_text = f"{doc_number} от {sign_date}"
        else:
            doc_info_text = f"от {sign_date}"

        # Сохраняем QR во временный файл
        qr_temp_path = None
        if line.qr:
            qr_binary_data = base64.b64decode(line.qr)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp_qr:
                tmp_qr.write(qr_binary_data)
                tmp_qr.flush()
                qr_temp_path = tmp_qr.name

        # Добавляем боковую подпись на КАЖДУЮ страницу
        for page_num, page in enumerate(reader.pages):
            if line.qr and qr_temp_path:
                page_box = page.mediabox
                page_width = float(page_box.width)
                page_height = float(page_box.height)

                overlay_io = io.BytesIO()
                c = canvas.Canvas(overlay_io, pagesize=(page_width, page_height))

                qr_size = 40
                margin_bottom = 14
                margin_left = 42

                c.saveState()
                c.translate(margin_left, margin_bottom + qr_size)
                c.rotate(90)

                c.drawImage(qr_temp_path, 0, 0, width=qr_size, height=qr_size)

                c.setFont(font_name, 7)
                c.drawString(qr_size + 5, 25, f"{signer_name} подписал(а)")
                c.drawString(qr_size + 5, 15, doc_info_text)

                c.restoreState()
                c.save()

                overlay_io.seek(0)
                overlay_reader = PdfReader(overlay_io)
                page.merge_page(overlay_reader.pages[0])

            writer.add_page(page)

        # Страница "ДОКУМЕНТ ПОДПИСАН"
        cert_page_io = io.BytesIO()
        c = canvas.Canvas(cert_page_io, pagesize=A4)
        page_width, page_height = A4

        c.setFont(font_name_bold, 14)
        c.drawCentredString(page_width / 2, page_height - 72, "ДОКУМЕНТ ПОДПИСАН")

        c.setFont(font_name, 12)
        c.drawString(72, page_height - 100, signer_name)

        y_position = page_height - 130
        line_height = 16

        cert_fields = [
            ("Дата выдачи сертификата:", str(line.certificate_start_date or '')),
            ("Дата окончания сертификата:", str(line.certificate_end_date or '')),
            ("Серийный номер:", str(line.serial_number or '')),
            ("Кем выдан:", str(line.issued_by or '')),
            ("Кому выдан:", str(line.issued_to or '')),
            ("Статус сертификата:", str(line.certificate_status or '')),
            ("Дата подписи:", sign_date),
        ]

        max_value_width = page_width - 72 - 200

        for label, value in cert_fields:
            c.setFont(font_name, 10)
            c.drawString(72, y_position, label)

            value_width = c.stringWidth(value, font_name, 10)
            if value_width > max_value_width:
                words = value.replace(',', ', ').split(' ')
                lines_list = []
                current_line = ''
                for word in words:
                    test_line = current_line + (' ' if current_line else '') + word
                    if c.stringWidth(test_line, font_name, 10) <= max_value_width:
                        current_line = test_line
                    else:
                        if current_line:
                            lines_list.append(current_line)
                        current_line = word
                if current_line:
                    lines_list.append(current_line)

                for i, text_line in enumerate(lines_list):
                    c.drawRightString(page_width - 72, y_position - (i * line_height), text_line)
                y_position -= line_height * len(lines_list)
            else:
                c.drawRightString(page_width - 72, y_position, value)
                y_position -= line_height

        if qr_temp_path:
            c.drawImage(qr_temp_path, page_width - 172, y_position - 100, width=100, height=100)

        c.showPage()
        c.save()
        cert_page_io.seek(0)
        cert_page_reader = PdfReader(cert_page_io)
        writer.add_page(cert_page_reader.pages[0])

        # Сохраняем итоговый PDF
        edited_pdf_io = io.BytesIO()
        writer.write(edited_pdf_io)
        edited_pdf_io.seek(0)
        edited_pdf_base64 = base64.b64encode(edited_pdf_io.read())

        new_attachment = self.env['ir.attachment'].sudo().create({
            'name': 'Подписан_' + attachment.name,
            'datas': edited_pdf_base64,
            'mimetype': 'application/pdf',
            'res_model': self._name,
            'res_id': self.id,
        })

        setattr(self.sudo(), field_name, [(3, attachment.id)])
        setattr(self.sudo(), field_name, [(4, new_attachment.id)])

        if qr_temp_path and os.path.exists(qr_temp_path):
            os.unlink(qr_temp_path)

    # ==================================================================
    # Вспомогательные методы
    # ==================================================================

    def additional_filter(self, init_approvers=None):
        return init_approvers

    def filter_connection(self, approval_groups=None):
        return approval_groups

    def check_group(self, group_ids):
        return group_ids.users

    def action_notify(self, notif_type, approver_id=None, reason=None):
        """Отправка уведомлений."""
        for record in self:
            user, template = record.create_uid, False
            template = self.env.ref(
                "correspondence.state_mixin_mail_template",
                raise_if_not_found=False
            )
            if not template:
                return
            template = template.sudo()
            template.model_id = (
                self.env["ir.model"]
                .sudo()
                .search([("model", "=", self._name)])
                .id
            )

            if notif_type == "agreement":
                user = approver_id
            elif notif_type in ["rework", "done", "approved"]:
                user = record.create_uid

            if user:
                try:
                    if user.notification_type == "email":
                        template.with_context(
                            object=record,
                            notif_type=notif_type,
                            reason=reason,
                            user=user,
                        ).send_mail(record.id, force_send=True)
                    else:
                        record.with_context(
                            object=record,
                            notif_type=notif_type,
                            reason=reason,
                            user=user,
                        ).message_post_with_template(template.id)
                except Exception as e:
                    _logger.warning(
                        "Ошибка отправки уведомления пользователю %s: %s",
                        user.name, str(e)
                    )
