# -*- coding: utf-8 -*-
"""Подстановка значений в загруженный .docx перед подписанием.

Генерируемые письма (направления, объяснительная, справка…) собираются
отчётом через docx_report_pro и подстановки не требуют. А письма, которые
сотрудник готовит в Word сам — обычное, гарантийное, уведомление об
изменении условий, — приходят с полями {{...}} в тексте, и без подстановки
эти поля уходят под ЭЦП как есть.

Движок — общий с модулем document_flow (models/doc_template_mixin.py).
Расходится с ним только точка подстановки, и то с оговоркой: там она
выбрана из тех же соображений, что здесь, — см. пункт 5. Когда движок
вынесут в appstream_approval, здесь останутся справочник полей и
_template_values.

Пять вещей, на которых здесь легко ошибиться:

1. Word рвёт поле на куски. {{date}} почти всегда лежит как {{da + te}} в
   соседних прогонах: автор поправил опечатку внутри, проверка орфографии
   разделила слово. Поиск по тексту отдельного прогона не находит ничего.
   Поэтому текст абзаца склеивается, поиск идёт по нему, а замена
   раскладывается обратно по прогонам: значение попадает в первый задетый,
   у остальных вырезается их часть. Заменить абзац целиком нельзя — в
   письмах половина строк с жирным, оформление потеряется.

2. document.paragraphs видит не всё: ни ячеек таблиц, ни надписей, ни
   колонтитулов, а шапка письма с номером обычно как раз в колонтитуле.
   Обходим XML каждой части пакета — сам документ и все колонтитулы,
   включая первую и чётные страницы.

   По той же причине обход идёт по элементам w:t, а не по
   paragraph.runs: runs отдаёт только прямых детей w:p, и поле внутри
   гиперссылки (w:hyperlink) или правки в режиме рецензирования (w:ins)
   остаётся незаполненным. Письмо, которое правили с включёнными
   исправлениями, — обычное дело.

3. Регистр не важен: {{Number}} и {{NUMBER}} — то же поле. Word
   автокапитализирует первую букву после точки, так что {{Number}} в
   письмах появляется сам.

4. python-docx пишет нестабильные байты: отметка времени в zip берётся из
   времени процесса, и то же содержимое после перезапуска даёт другой хеш.
   Фреймворк сверяет хеш документа из предпросмотра с хешем при подписании
   и ругается «Документ изменился после предпросмотра». Поэтому zip
   пересобирается с постоянной датой 1980-01-01.

5. Подстановка делается при сборке файла на подпись
   (_esp_field_sources), а не после неё. Этапы Утверждения и Мед.
   освидетельствования работают на «XML с хешем документа» (esp_type =
   xml_doc, выставлен migrations/19.0.1.0.1): хеш файла входит в подпись,
   а вложение помечается signed_with_esp сразу при сохранении подписи.
   Позже менять в файле нечего — байты заморожены. Поэтому номер в
   подстановку идёт предварительный (name_draft), а после подписи
   сверяется с присвоенным: см. _assign_document_number.

   Файл в поле записи при этом остаётся шаблоном: фреймворк копирует
   байты в отдельное вложение документа на подпись
   (_create_documents), так что после доработки письмо можно отправить
   на согласование ещё раз.

В PDF подстановка не делается: текст там набор глифов с координатами, и
замена «ИСХ-2026-0147» вместо «{{number}}» поехала бы по вёрстке. PDF
только проверяется, чтобы шаблон не ушёл под подпись незаполненным.
"""

import base64
import io
import logging
import re
import zipfile

from odoo import _, models
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)

# {{ поле }} — пробелы необязательны, точка допускается, чтобы опечатка в
# составном имени попала в проверку, а не осталась в тексте.
FIELD = re.compile(r"\{\{\s*([A-Za-z_][\w.]*)\s*\}\}")

DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# Фиксированная отметка времени в zip: см. пункт 4 в шапке файла.
STABLE_DATE = (1980, 1, 1, 0, 0, 0)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ----------------------------------------------------------------------
# Чистые функции: работают с байтами файла и словарём значений
# ----------------------------------------------------------------------

def _header_footer_roots(document):
    """Корневые элементы колонтитулов, которые в документе действительно есть.

    Берутся по связям основной части пакета, а не через section.header и
    соседей: у python-docx эти свойства при обращении к _element СОЗДАЮТ
    колонтитул, если у секции его нет (_get_or_add_definition). Обход
    дописывал в каждое письмо шесть пустых частей header/footer — уже в
    файл, который уходит под подпись, — а пустой верхний колонтитул на
    бланке с узким верхним полем сдвигает текст письма вниз.

    Связи части уникальны, поэтому общий для нескольких секций колонтитул
    встречается здесь один раз.
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    roots, seen = [], set()
    for rel in document.part.rels.values():
        if rel.is_external or rel.reltype not in (RT.HEADER, RT.FOOTER):
            continue
        part = rel.target_part
        if id(part) not in seen:
            seen.add(id(part))
            roots.append(part.element)
    return roots


def _iter_paragraph_elements(document):
    """Все абзацы пакета: тело, таблицы, надписи и колонтитулы."""
    for root in [document.element.body] + _header_footer_roots(document):
        for paragraph in root.iter(_W + "p"):
            yield paragraph


def _paragraph_text(paragraph):
    return "".join(run.text or "" for run in paragraph.iter(_W + "t"))


def _fill_paragraph(paragraph, values, missing):
    """Подставить значения в один абзац. Возвращает число замен.

    Прогоны не склеиваются: значение кладётся в первый задетый, у остальных
    вырезается только их часть поля, поэтому жирный, курсив и размер
    остального текста остаются как были.

    Незнакомое поле остаётся в тексте как есть и попадает в missing.
    Затирать его пустой строкой нельзя: после ЭЦП лучше увидеть
    {{otpravitel}}, чем молча потерять строку, — а проверка перед
    отправкой (check_templates) до этого и не должна допускать.
    """
    runs = list(paragraph.iter(_W + "t"))
    if not runs:
        return 0
    texts = [run.text or "" for run in runs]
    full = "".join(texts)
    if "{{" not in full:
        return 0

    spans = []
    for match in FIELD.finditer(full):
        name = match.group(1).lower()
        if name not in values:
            missing.add(match.group(1))
            continue
        spans.append((match.start(), match.end(), "%s" % (values[name] or "")))
    if not spans:
        return 0

    # Что эмитить на каждой позиции сквозного текста: значение — на первом
    # символе поля, пусто — на остальных его символах.
    pieces = list(full)
    for start, end, replacement in spans:
        pieces[start] = replacement
        for index in range(start + 1, end):
            pieces[index] = ""

    offset = 0
    for run, text in zip(runs, texts):
        run.text = "".join(pieces[offset:offset + len(text)])
        offset += len(text)
    return len(spans)


def _stable_docx(data):
    """Пересобрать zip с постоянной отметкой времени: см. пункт 4 в шапке.

    Порядок частей сохраняется — [Content_Types].xml должен остаться первым.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=STABLE_DATE)
            info.compress_type = item.compress_type
            info.external_attr = item.external_attr
            info.create_system = item.create_system
            target.writestr(info, source.read(item.filename))
    return buffer.getvalue()


def scan_docx(data):
    """Имена полей {{...}}, встречающиеся в документе."""
    from docx import Document
    document = Document(io.BytesIO(data))
    found = set()
    for paragraph in _iter_paragraph_elements(document):
        found.update(FIELD.findall(_paragraph_text(paragraph)))
    return found


def scan_pdf(data):
    """Имена полей {{...}} в тексте PDF.

    Через pypdf, а где его нет — PyPDF2. Сейчас PyPDF2 есть всегда: без
    него не загружается сам appstream_approval (esp_copy.py, печатные
    версии). pypdf — его поддерживаемый преемник, его ставит Odoo 19 на
    Python 3.13; первая попытка — задел на переход фреймворка. Сканировать
    байты файла бессмысленно: содержимое страниц лежит в сжатых потоках.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            _logger.info("Ни pypdf, ни PyPDF2 не установлены — PDF на поля "
                         "подстановки не проверяется")
            return set()
    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        # PDF мог быть сканом, с защитой или битым — не повод не пустить письмо
        _logger.exception("Не удалось прочитать текст PDF для проверки полей")
        return set()
    return set(FIELD.findall(text))


def fill_docx(data, values):
    """Подставить значения. Возвращает (байты, число замен, неизвестные поля).

    Байты возвращаются только если что-то заменилось: иначе файл остаётся
    прежним и его хеш не меняется.
    """
    from docx import Document
    document = Document(io.BytesIO(data))
    lowered = {name.lower(): value for name, value in values.items()}
    missing, replaced = set(), 0
    for paragraph in _iter_paragraph_elements(document):
        replaced += _fill_paragraph(paragraph, lowered, missing)
    if not replaced:
        return None, 0, missing
    buffer = io.BytesIO()
    document.save(buffer)
    return _stable_docx(buffer.getvalue()), replaced, missing


# ----------------------------------------------------------------------
# Миксин для модели документа
# ----------------------------------------------------------------------

class DocxTemplateMixin(models.AbstractModel):
    _name = "corr.docx.template.mixin"
    _description = "Подстановка значений в загруженный .docx"

    # Справочник полей подстановки: (имя, что подставляется, пример).
    # Заполняется в модели документа — из него строится подсказка
    # исполнителю. Проверка шаблона им не ограничивается: см.
    # _template_known_names.
    TEMPLATE_FIELDS = []

    # Второе имя для того же значения: {алиас: настоящее имя}. Заполняется
    # в модели документа, которая и подставляет их в _template_values.
    TEMPLATE_ALIASES = {}

    def _template_values(self):
        """Значения для подстановки: {имя поля: строка}. Переопределяется."""
        self.ensure_one()
        return {}

    def _template_known_names(self):
        """Имена, которые система умеет заполнять.

        Объединение справочника и ключей _template_values: справочник
        нужен для подсказки с описанием и примером, а ключи — чтобы
        проверка не отвергла поле, которое заполняется, но в справочник
        его забыли внести.
        """
        self.ensure_one()
        names = {name.lower() for name, _description, _example in self.TEMPLATE_FIELDS}
        return names | {name.lower() for name in self._template_values()}

    def _template_sources(self):
        """Поля записи с файлами, в которых делается подстановка."""
        return []

    def _template_attachments(self):
        """Вложения из полей _template_sources()."""
        self.ensure_one()
        attachments = self.env["ir.attachment"]
        for name in self._template_sources():
            if name in self._fields:
                attachments |= self[name]
        return attachments

    @staticmethod
    def _is_docx(attachment):
        return (attachment.mimetype == DOCX_MIMETYPE
                or (attachment.name or "").lower().endswith(".docx"))

    @staticmethod
    def _is_pdf(attachment):
        return (attachment.mimetype == "application/pdf"
                or (attachment.name or "").lower().endswith(".pdf"))

    def check_templates(self):
        """Не пустить документ дальше, если в шаблоне есть незнакомое поле.

        Без этой проверки опечатка вроде {{otpravitel}} доедет до подписи и
        закрепится ЭЦП навсегда. PDF тоже проверяется: подставить в него
        нельзя, но и уйти под подпись с {{...}} он не должен.
        """
        for record in self:
            known = record._template_known_names()
            for attachment in record._template_attachments().sudo():
                if not attachment.datas:
                    continue
                if record._is_docx(attachment):
                    found = scan_docx(base64.b64decode(attachment.datas))
                elif record._is_pdf(attachment):
                    found = scan_pdf(base64.b64decode(attachment.datas))
                    if found:
                        raise ValidationError(_(
                            "В файле «%(file)s» есть поля подстановки "
                            "%(found)s, но это PDF — переписать текст в нём, "
                            "не сломав вёрстку, нельзя.\n\n"
                            "Приложите письмо в формате .docx.",
                            file=attachment.name,
                            found=", ".join("{{%s}}" % name for name in sorted(found)),
                        ))
                    continue
                else:
                    continue
                unknown = sorted(
                    name for name in found if name.lower() not in known)
                if not unknown:
                    continue
                raise ValidationError(_(
                    "В файле «%(file)s» есть поля, которых система не знает: "
                    "%(unknown)s.\n\nДоступные поля: %(known)s",
                    file=attachment.name,
                    unknown=", ".join("{{%s}}" % name for name in unknown),
                    known=", ".join("{{%s}}" % name for name in sorted(known)),
                ))
        return True

    def fill_template(self, raw, filename):
        """Подставить значения в байты документа. None — менять нечего.

        Сбой подстановки не отменяет подпись, но и молчать о нём нельзя:
        незаполненное поле видно уже после ЭЦП. Поэтому — в чаттер.
        """
        self.ensure_one()
        if not (filename or "").lower().endswith(".docx"):
            return None
        try:
            filled, replaced, missing = fill_docx(raw, self._template_values())
        except Exception as error:
            _logger.exception("Подстановка в %s не выполнена (%s)", filename, self)
            self._notify_template_problem(_(
                "Не удалось подставить данные в «%(file)s»: %(error)s\n"
                "В подписанный файл уходят незаполненные поля.",
                file=filename, error=error))
            return None
        if missing:
            # В чаттер не пишем: _esp_field_sources вызывается и на
            # предпросмотре, сообщение повторялось бы на каждый просмотр.
            # Незнакомое поле до подписи и не должно доходить — его ловит
            # check_templates при отправке на согласование; если оно
            # всё-таки здесь, это расхождение в коде, а не ошибка автора.
            _logger.warning("В %s остались незнакомые поля: %s",
                            filename, ", ".join(sorted(missing)))
        return filled

    def _notify_template_problem(self, body):
        """Сообщить о сбое подстановки в чаттер, а не только в лог.

        Сбой гасится намеренно — подпись важнее заполнения, — но тихо
        отдавать на подпись письмо с {{number}} нельзя.
        """
        self.ensure_one()
        if hasattr(self, "message_post"):
            try:
                self.sudo().message_post(body=body)
            except Exception:
                _logger.exception("Не удалось записать в чаттер сбой подстановки")

    def action_show_template_fields(self):
        """Подсказка исполнителю: какие поля можно вставить в письмо.

        Отдельного визарда не заводим — список статический.
        """
        self.ensure_one()
        described = {name.lower() for name, _d, _e in self.TEMPLATE_FIELDS}
        lines = ["{{%s}} — %s (например: %s)" % (name, description, example)
                 for name, description, example in self.TEMPLATE_FIELDS]
        aliases = self.TEMPLATE_ALIASES
        if aliases:
            lines.append("")
            lines += ["{{%s}} — то же, что {{%s}}" % (alias, aliases[alias])
                      for alias in sorted(aliases)]
            described |= {alias.lower() for alias in aliases}
        # поля, которые заполняются, но в справочник не внесены
        extra = [name for name in sorted(self._template_values())
                 if name.lower() not in described]
        if extra:
            lines.append("")
            lines += ["{{%s}}" % name for name in extra]
        # UserError, а не ValidationError: это подсказка, а не отказ в
        # сохранении, — Odoo покажет её диалогом с сохранением переносов.
        raise UserError(_(
            "Поля подстановки для исходящего письма.\n\n"
            "Вставьте их в .docx обычным текстом — при сборке документа на "
            "подпись система заменит их данными письма. Регистр не важен, "
            "поля в колонтитулах и таблицах тоже заполняются.\n\n%s",
            "\n".join(lines)))
